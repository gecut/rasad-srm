#!/usr/bin/env python3
"""
extract-students-from-xlsx.py
Extracts and normalizes student records directly from original Excel spreadsheets:
- /Users/mm25zamanian/Downloads/پشت خطی ها.xlsx (Sheet 'نهایی')
- Matches class enrolments from data/classes.json (derived from کلاس های دبستان.xlsx)
Generates clean data/students.json without any CSV column-shift corruption.
"""

import os
import sys
import json
import re
import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, '..'))
EXCEL_PATH = '/Users/mm25zamanian/Downloads/پشت خطی ها.xlsx'
CLASSES_JSON_PATH = os.path.join(ROOT_DIR, 'data', 'classes.json')
OUTPUT_JSON_PATH = os.path.join(ROOT_DIR, 'data', 'students.json')

def to_latin_digits(s):
    if not s:
        return ''
    fa = '۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩'
    en = '01234567890123456789'
    return str(s).translate(str.maketrans(fa, en))

def clean_str(val):
    if val is None:
        return None
    s = str(val).strip()
    s = s.replace('ي', 'ی').replace('ك', 'ک')
    s = re.sub(r'[\u200c\u200e\u200f]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s if s else None

def norm_name(text):
    if not text:
        return ''
    s = clean_str(text)
    return s.replace(' ', '')

def main():
    print(f"Loading Excel file: {EXCEL_PATH} ...")
    if not os.path.exists(EXCEL_PATH):
        print(f"Error: {EXCEL_PATH} not found!")
        sys.exit(1)

    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    if 'نهایی' not in wb.sheetnames:
        print(f"Error: Sheet 'نهایی' not found in {EXCEL_PATH}! Available: {wb.sheetnames}")
        sys.exit(1)

    sheet = wb['نهایی']
    rows = list(sheet.iter_rows(values_only=True))
    print(f"Total rows in sheet 'نهایی': {len(rows)}")

    # Load classes
    student_to_class = {}
    if os.path.exists(CLASSES_JSON_PATH):
        with open(CLASSES_JSON_PATH, 'r', encoding='utf-8') as f:
            classes_data = json.load(f)
        for c in classes_data:
            for sname in c.get('studentNames', []):
                student_to_class[norm_name(sname)] = c['title']
        print(f"Indexed {len(student_to_class)} student class assignments from {CLASSES_JSON_PATH}.")

    students = []
    stats = {
        'total': 0,
        'waitlisted': 0,
        'class_seeker': 0,
        'referred_to_teacher': 0,
        'unknown_grade': 0,
        'missing_last_name': 0,
        'has_father_mobile': 0,
        'has_mother_mobile': 0,
        'has_landline': 0,
        'has_address': 0,
        'has_ceremonies': 0,
    }

    for idx, r in enumerate(rows[1:], start=2):
        # Empty row check
        if not any(r):
            continue

        raw_first = clean_str(r[0])
        raw_last = clean_str(r[1])
        if not raw_first and not raw_last:
            continue

        # Grade
        grade_val = r[2]
        grade = None
        if grade_val is not None:
            try:
                g = int(float(to_latin_digits(str(grade_val)).strip()))
                if 1 <= g <= 6:
                    grade = g
            except (ValueError, TypeError):
                grade = None

        firstName = raw_first or '(نامشخص)'
        lastName = raw_last or '(ثبت‌نشده)'
        if not raw_last:
            stats['missing_last_name'] += 1

        # Phones
        raw_landline = r[3]
        raw_father = r[4]
        raw_mother = r[5]

        father_mobile = None
        mother_mobile = None
        student_mobile = None
        landline = None
        phone_notes = []

        # 1. Father mobile
        if raw_father is not None:
            s_fat = to_latin_digits(str(raw_father)).strip()
            if s_fat.endswith('.0'):
                s_fat = s_fat[:-2]
            m = re.search(r'(09\d{9})|(\+989\d{9})|(9\d{9})', s_fat)
            if m:
                matched = m.group(0)
                if matched.startswith('+98'):
                    matched = '0' + matched[3:]
                elif len(matched) == 10 and matched.startswith('9'):
                    matched = '0' + matched
                father_mobile = matched
                # Check for extra remarks attached to the number
                remaining = re.sub(r'[\s0-9()\-+]+', '', s_fat).strip()
                if remaining:
                    phone_notes.append(f'توضیح شماره پدر: {remaining}')
            else:
                phone_notes.append(f'شماره همراه پدر در اکسل: {raw_father}')

        # 2. Mother mobile
        if raw_mother is not None:
            s_mot = to_latin_digits(str(raw_mother)).strip()
            if s_mot.endswith('.0'):
                s_mot = s_mot[:-2]
            m = re.search(r'(09\d{9})|(\+989\d{9})|(9\d{9})', s_mot)
            if m:
                matched = m.group(0)
                if matched.startswith('+98'):
                    matched = '0' + matched[3:]
                elif len(matched) == 10 and matched.startswith('9'):
                    matched = '0' + matched
                mother_mobile = matched
                remaining = re.sub(r'[\s0-9()\-+]+', '', s_mot).strip()
                if remaining:
                    phone_notes.append(f'توضیح شماره مادر: {remaining}')
            else:
                # Check if it's a landline (e.g. 5138798733)
                m_land = re.search(r'(0?51\d{8})|(\d{8})', s_mot)
                if m_land:
                    land_num = m_land.group(0)
                    if not landline:
                        landline = land_num
                    else:
                        phone_notes.append(f'شماره مادر در اکسل: {raw_mother}')
                else:
                    phone_notes.append(f'شماره همراه مادر در اکسل: {raw_mother}')

        # 3. Landline / Home phone
        if raw_landline is not None:
            s_land = to_latin_digits(str(raw_landline)).strip()
            if s_land.endswith('.0'):
                s_land = s_land[:-2]
            # Check if it's a mobile
            m_mob = re.search(r'(09\d{9})|(9\d{9})', s_land)
            if m_mob:
                mob = m_mob.group(0)
                if len(mob) == 10:
                    mob = '0' + mob
                if not student_mobile:
                    student_mobile = mob
                elif not father_mobile:
                    father_mobile = mob
                elif not mother_mobile:
                    mother_mobile = mob
                else:
                    phone_notes.append(f'موبایل ثبت شده در تلفن منزل: {mob}')
            else:
                m_land = re.search(r'(\d{7,11})', s_land)
                if m_land:
                    landline = m_land.group(0)
                else:
                    phone_notes.append(f'تلفن منزل در اکسل: {raw_landline}')

        # Neighborhood
        raw_neigh = clean_str(r[6])
        neighborhood = raw_neigh
        if neighborhood == 'وکیل اباد راست':
            neighborhood = 'وکیل آباد راست'

        # Address
        address = clean_str(r[7])

        # Referrer
        referrer = clean_str(r[8])

        # Ceremonies
        # Col 11: نیمه1402, Col 12: غدیر1403, Col 13: نیمه1403, Col 14: غدیر1404, Col 15: غدیر1405 (mapped to نیمه1404)
        ceremony_map = {}
        ceremony_keys = [
            (11, 'نیمه1402'),
            (12, 'غدیر1403'),
            (13, 'نیمه1403'),
            (14, 'غدیر1404'),
            (15, 'نیمه1404'),  # غدیر1405 in spreadsheet is actually نیمه1404 (جشن نیمه شعبان ۱۴۰۴)
        ]
        for col_idx, ckey in ceremony_keys:
            if col_idx < len(r):
                val = r[col_idx]
                if val is not None and str(val).strip():
                    try:
                        s_val = str(int(float(str(val).strip())))
                    except (ValueError, TypeError):
                        s_val = str(val).strip()
                    ceremony_map[ckey] = s_val

        # Class matching
        full_name_norm = norm_name(f'{firstName}{lastName}')
        assigned_class = student_to_class.get(full_name_norm, None)

        # Status rules per business rules and user instructions:
        # 1. Grade < 4 (1, 2, 3): readinessStatus = 'waitlisted', lifecycleStatus = 'unknown'
        # 2. Grade >= 4 (4, 5, 6):
        #    - If matched to active class: readinessStatus = 'normal', lifecycleStatus = 'referred_to_teacher'
        #    - If not in class: readinessStatus = 'normal', lifecycleStatus = 'class_seeker'
        # 3. Grade is None: readinessStatus = 'normal', lifecycleStatus = 'unknown'
        if grade is None:
            readinessStatus = 'normal'
            lifecycleStatus = 'unknown'
            stats['unknown_grade'] += 1
        elif grade < 4:
            readinessStatus = 'waitlisted'
            lifecycleStatus = 'unknown'
            stats['waitlisted'] += 1
        else:
            readinessStatus = 'normal'
            if assigned_class:
                lifecycleStatus = 'referred_to_teacher'
                stats['referred_to_teacher'] += 1
            else:
                lifecycleStatus = 'class_seeker'
                stats['class_seeker'] += 1

        # Notes
        raw_notes = clean_str(r[16]) if len(r) > 16 else None
        notes_list = []
        if raw_notes:
            notes_list.append(raw_notes)
        if phone_notes:
            notes_list.extend(phone_notes)
        notes = ' | '.join(notes_list) if notes_list else ''

        # Statistics
        stats['total'] += 1
        if father_mobile:
            stats['has_father_mobile'] += 1
        if mother_mobile:
            stats['has_mother_mobile'] += 1
        if landline:
            stats['has_landline'] += 1
        if address:
            stats['has_address'] += 1
        if ceremony_map:
            stats['has_ceremonies'] += 1

        st = {
            'id': f'row_{idx}',
            'firstName': firstName,
            'lastName': lastName,
            'grade': grade,
            'origin': 'import',
            'readinessStatus': readinessStatus,
            'lifecycleStatus': lifecycleStatus,
            'currentClass': assigned_class,
            'fatherMobile': father_mobile,
            'motherMobile': mother_mobile,
            'mobile': student_mobile,
            'landline': landline,
            'neighborhood': neighborhood,
            'address': address,
            'referrer': referrer,
            'notes': notes,
            'ceremonies': ceremony_map,
        }
        students.append(st)

    # Write output to JSON
    with open(OUTPUT_JSON_PATH, 'w', encoding='utf-8') as f:
        json.dump(students, f, ensure_ascii=False, indent=2)

    print(f"\nSuccessfully wrote {len(students)} student records to {OUTPUT_JSON_PATH}!")
    print("\nExtraction Summary:")
    print(json.dumps(stats, indent=2, ensure_ascii=False))

if __name__ == '__main__':
    main()
